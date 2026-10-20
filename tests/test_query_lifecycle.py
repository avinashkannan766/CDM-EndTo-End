
import sys
import unittest
from datetime import date
from pathlib import Path

from openpyxl import Workbook

PROJECT_DIR = Path(__file__).resolve().parent.parent
SCRIPTS_DIR = PROJECT_DIR / "scripts"
sys.path.insert(0, str(SCRIPTS_DIR))

from manage_query_lifecycle import transition_query


HEADERS = [
    "Query_ID",
    "Query_Status",
    "Response",
    "Resolution",
    "Date_Closed",
    "Reviewer_Notes",
]

COLUMN_MAP = {
    name: index
    for index, name in enumerate(HEADERS, start=1)
}


def make_query(status="Open", response=None, resolution=None):
    """Create an isolated workbook containing one training query."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Generated Queries"
    sheet.append(HEADERS)
    sheet.append([
        "QRY-TEST-001",
        status,
        response,
        resolution,
        None,
        None,
    ])
    return workbook, sheet


class TestQueryLifecycle(unittest.TestCase):

    def test_valid_transition_to_answered(self):
        workbook, sheet = make_query()

        try:
            transition_query(
                sheet,
                COLUMN_MAP,
                "QRY-TEST-001",
                "Answered",
                response="Training response",
            )

            self.assertEqual(sheet.cell(2, 2).value, "Answered")
            self.assertEqual(
                sheet.cell(2, 3).value,
                "Training response",
            )
        finally:
            workbook.close()

    def test_cannot_answer_without_response(self):
        workbook, sheet = make_query()

        try:
            with self.assertRaisesRegex(ValueError, "Response"):
                transition_query(
                    sheet,
                    COLUMN_MAP,
                    "QRY-TEST-001",
                    "Answered",
                    response="   ",
                )
        finally:
            workbook.close()

    def test_cannot_resolve_without_resolution(self):
        workbook, sheet = make_query(
            status="Answered",
            response="Existing training response",
        )
        original_response = sheet.cell(2, 3).value

        try:
            with self.assertRaisesRegex(ValueError, "Resolution"):
                transition_query(
                    sheet,
                    COLUMN_MAP,
                    "QRY-TEST-001",
                    "Resolved",
                    response="Updated training response",
                )

            # A rejected transition must not partially change the query.
            self.assertEqual(sheet.cell(2, 2).value, "Answered")
            self.assertEqual(sheet.cell(2, 3).value, original_response)
            self.assertIsNone(sheet.cell(2, 4).value)
        finally:
            workbook.close()

    def test_cannot_skip_directly_from_open_to_closed(self):
        workbook, sheet = make_query()

        try:
            with self.assertRaisesRegex(ValueError, "not allowed"):
                transition_query(
                    sheet,
                    COLUMN_MAP,
                    "QRY-TEST-001",
                    "Closed",
                )

            self.assertEqual(sheet.cell(2, 2).value, "Open")
            self.assertIsNone(sheet.cell(2, 5).value)
        finally:
            workbook.close()

    def test_unknown_query_id_is_rejected(self):
        workbook, sheet = make_query()

        try:
            with self.assertRaisesRegex(ValueError, "not found"):
                transition_query(
                    sheet,
                    COLUMN_MAP,
                    "QRY-UNKNOWN",
                    "Answered",
                    response="Training response",
                )
        finally:
            workbook.close()

    def test_invalid_status_is_rejected(self):
        workbook, sheet = make_query()

        try:
            with self.assertRaisesRegex(ValueError, "Invalid status"):
                transition_query(
                    sheet,
                    COLUMN_MAP,
                    "QRY-TEST-001",
                    "Pending",
                )

            self.assertEqual(sheet.cell(2, 2).value, "Open")
        finally:
            workbook.close()

    def test_complete_lifecycle_and_closure_date(self):
        workbook, sheet = make_query()

        try:
            transition_query(
                sheet,
                COLUMN_MAP,
                "QRY-TEST-001",
                "Answered",
                response="Training response",
            )
            transition_query(
                sheet,
                COLUMN_MAP,
                "QRY-TEST-001",
                "Resolved",
                resolution="Training resolution",
            )
            transition_query(
                sheet,
                COLUMN_MAP,
                "QRY-TEST-001",
                "Closed",
            )

            self.assertEqual(sheet.cell(2, 2).value, "Closed")
            self.assertEqual(
                sheet.cell(2, 3).value,
                "Training response",
            )
            self.assertEqual(
                sheet.cell(2, 4).value,
                "Training resolution",
            )
            self.assertEqual(sheet.cell(2, 5).value, date.today())
        finally:
            workbook.close()

    def test_closed_query_cannot_be_reopened_directly_to_answered(self):
        workbook, sheet = make_query(
            status="Closed",
            response="Training response",
            resolution="Training resolution",
        )

        try:
            with self.assertRaisesRegex(ValueError, "not allowed"):
                transition_query(
                    sheet,
                    COLUMN_MAP,
                    "QRY-TEST-001",
                    "Answered",
                )

            self.assertEqual(sheet.cell(2, 2).value, "Closed")
        finally:
            workbook.close()


if __name__ == "__main__":
    unittest.main(verbosity=2)
