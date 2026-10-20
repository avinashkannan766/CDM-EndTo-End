
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

PROJECT_DIR = Path(__file__).resolve().parent.parent
SCRIPTS_DIR = PROJECT_DIR / "scripts"
sys.path.insert(0, str(SCRIPTS_DIR))

import create_correction_audit_log as audit


class TestCorrectionAudit(unittest.TestCase):

    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.temp_path = Path(self.temp_dir.name)

        self.source_file = self.temp_path / "query_lifecycle_working.xlsx"
        self.output_file = self.temp_path / "data_correction_audit_log.xlsx"

        # Create an isolated source workbook with two fictional queries.
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Generated Queries"

        sheet.append([
            "Query_ID",
            "Subject_ID",
            "Form_Domain",
            "Record_Reference",
            "Query_Status",
        ])

        sheet.append([
            "QRY-TEST-001",
            "DM-TEST-001",
            "Laboratory",
            "LAB-TEST-001",
            "Open",
        ])

        sheet.append([
            "QRY-TEST-002",
            "DM-TEST-002",
            "Visits",
            "VISIT-TEST-002",
            "Resolved",
        ])

        workbook.save(self.source_file)
        workbook.close()

        self.query_file_patch = patch.object(
            audit, "QUERY_FILE", self.source_file
        )
        self.output_file_patch = patch.object(
            audit, "OUTPUT_FILE", self.output_file
        )

        self.query_file_patch.start()
        self.output_file_patch.start()

    def tearDown(self):
        self.query_file_patch.stop()
        self.output_file_patch.stop()
        self.temp_dir.cleanup()

    def test_loads_query_references(self):
        references = audit.load_query_references()

        self.assertEqual(len(references), 2)
        self.assertEqual(references[0][0], "QRY-TEST-001")
        self.assertEqual(references[1][0], "QRY-TEST-002")
        self.assertEqual(references[0][5], "Pending review")

    def test_creates_expected_sheets_and_columns(self):
        audit.create_audit_workbook()

        workbook = load_workbook(self.output_file, data_only=True)

        try:
            self.assertEqual(
                workbook.sheetnames,
                ["Correction Audit", "Query Reference", "Instructions"],
            )

            audit_sheet = workbook["Correction Audit"]
            headers = [
                cell.value for cell in audit_sheet[1]
            ]

            self.assertEqual(headers, audit.AUDIT_COLUMNS)
            self.assertEqual(audit_sheet.max_column, 14)
        finally:
            workbook.close()

    def test_audit_template_does_not_claim_a_correction(self):
        audit.create_audit_workbook()

        workbook = load_workbook(self.output_file, data_only=True)

        try:
            sheet = workbook["Correction Audit"]

            self.assertEqual(sheet["J2"].value, "Pending Review")
            self.assertEqual(
                sheet["N2"].value,
                "No correction has been applied.",
            )

            # Proposed change fields must remain blank.
            for column in ["A", "B", "C", "D", "E", "F", "G", "H", "I"]:
                self.assertIsNone(sheet[f"{column}2"].value)

            self.assertIsNone(sheet["K2"].value)
            self.assertIsNone(sheet["L2"].value)
            self.assertIsNone(sheet["M2"].value)
        finally:
            workbook.close()

    def test_query_reference_sheet_contains_all_source_queries(self):
        audit.create_audit_workbook()

        workbook = load_workbook(self.output_file, data_only=True)

        try:
            sheet = workbook["Query Reference"]

            self.assertEqual(sheet.max_row, 3)
            self.assertEqual(sheet["A2"].value, "QRY-TEST-001")
            self.assertEqual(sheet["A3"].value, "QRY-TEST-002")
            self.assertEqual(sheet["F2"].value, "Pending review")
            self.assertEqual(sheet["F3"].value, "Pending review")
        finally:
            workbook.close()

    def test_refuses_to_overwrite_existing_output(self):
        self.output_file.write_bytes(b"existing file must remain unchanged")
        original_bytes = self.output_file.read_bytes()

        with self.assertRaises(FileExistsError):
            audit.create_audit_workbook()

        self.assertEqual(self.output_file.read_bytes(), original_bytes)

    def test_missing_source_workbook_is_reported(self):
        self.source_file.unlink()

        with self.assertRaises(FileNotFoundError):
            audit.load_query_references()

    def test_missing_required_source_column_is_reported(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Generated Queries"

        sheet.append([
            "Query_ID",
            "Subject_ID",
            "Form_Domain",
            # Record_Reference intentionally missing
            "Query_Status",
        ])

        sheet.append([
            "QRY-TEST-003",
            "DM-TEST-003",
            "Laboratory",
            "Open",
        ])

        workbook.save(self.source_file)
        workbook.close()

        with self.assertRaisesRegex(ValueError, "Record_Reference"):
            audit.load_query_references()


if __name__ == "__main__":
    unittest.main(verbosity=2)
