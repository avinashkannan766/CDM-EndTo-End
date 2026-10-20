
from pathlib import Path
from datetime import datetime
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


BASE_DIR = Path(__file__).resolve().parent.parent
REPORT_DIR = BASE_DIR / "reports"

CLEAN_FINDINGS_FILE = REPORT_DIR / "clean_edit_check_findings.xlsx"
TEST_FINDINGS_FILE = REPORT_DIR / "test_edit_check_findings.xlsx"
QUERY_FILE = REPORT_DIR / "generated_query_log.xlsx"
LIFECYCLE_FILE = REPORT_DIR / "query_lifecycle_working.xlsx"
RECONCILIATION_FILE = REPORT_DIR / "reconciliation_report.xlsx"

OUTPUT_FILE = REPORT_DIR / "cdm_summary_report.xlsx"


def read_sheet(file_path, sheet_name):
    """Read a worksheet into a list of dictionaries."""
    if not file_path.exists():
        raise FileNotFoundError(f"Required report not found: {file_path}")

    workbook = load_workbook(file_path, data_only=True, read_only=True)

    try:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(
                f"Sheet '{sheet_name}' not found in {file_path.name}. "
                f"Available sheets: {workbook.sheetnames}"
            )

        worksheet = workbook[sheet_name]
        rows = worksheet.iter_rows(values_only=True)
        headers = next(rows, None)

        if not headers:
            return []

        headers = [
            str(value).strip() if value is not None else ""
            for value in headers
        ]

        records = []

        for row in rows:
            if not any(value is not None for value in row):
                continue

            record = {
                headers[index]: row[index] if index < len(row) else None
                for index in range(len(headers))
                if headers[index]
            }
            records.append(record)

        return records

    finally:
        workbook.close()


def get_value(record, *possible_names, default=None):
    """Find a field while allowing for small header-name differences."""
    for name in possible_names:
        if name in record:
            return record[name]
    return default


def write_table(worksheet, headers, rows):
    """Write a formatted table to a worksheet."""
    worksheet.append(headers)

    for row in rows:
        worksheet.append([row.get(header) for header in headers])

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78"
    )

    for cell in worksheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True
        )

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions

    for column_cells in worksheet.columns:
        column_letter = get_column_letter(column_cells[0].column)
        maximum_length = max(
            len(str(cell.value)) if cell.value is not None else 0
            for cell in column_cells
        )
        worksheet.column_dimensions[column_letter].width = min(
            max(maximum_length + 2, 14),
            45
        )

    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True
            )


def main():
    # Read existing edit-check reports.
    clean_findings = read_sheet(
        CLEAN_FINDINGS_FILE, "Findings"
    )
    test_findings = read_sheet(
        TEST_FINDINGS_FILE, "Findings"
    )

    # Read generated queries and the demonstrated lifecycle workbook.
    queries = read_sheet(
        QUERY_FILE, "Generated Queries"
    )
    lifecycle_queries = read_sheet(
        LIFECYCLE_FILE, "Generated Queries"
    )

    # Read reconciliation results.
    reconciliation_summary = read_sheet(
        RECONCILIATION_FILE, "Summary"
    )
    reconciliation_findings = read_sheet(
        RECONCILIATION_FILE, "Findings"
    )

    # Summarize edit-check findings.
    test_errors = sum(
        1 for finding in test_findings
        if str(get_value(finding, "Severity", default="")).strip().lower()
        == "error"
    )

    test_reviews = sum(
        1 for finding in test_findings
        if str(get_value(finding, "Severity", default="")).strip().lower()
        in {"review", "review flag"}
    )

    # Count generated query statuses.
    query_status_counts = {}

    for query in queries:
        status = str(
            get_value(query, "Query_Status", "Status", default="Unknown")
        ).strip()

        query_status_counts[status] = (
            query_status_counts.get(status, 0) + 1
        )

    # Count statuses in the lifecycle demonstration separately.
    lifecycle_status_counts = {}

    for query in lifecycle_queries:
        status = str(
            get_value(query, "Query_Status", "Status", default="Unknown")
        ).strip()

        lifecycle_status_counts[status] = (
            lifecycle_status_counts.get(status, 0) + 1
        )

    # Convert reconciliation summary rows into a dictionary.
    reconciliation_metrics = {}

    for item in reconciliation_summary:
        metric = get_value(
            item, "Metric", "Check", "Measure", "Name"
        )
        value = get_value(
            item, "Value", "Count", "Result"
        )

        if metric is not None:
            reconciliation_metrics[str(metric).strip()] = value

    # Build the consolidated workbook.
    workbook = Workbook()

    overview = workbook.active
    overview.title = "CDM Overview"

    overview_rows = [
        {
            "Section": "Project",
            "Metric": "Study ID",
            "Result": "DM-001",
            "Interpretation": "Fictional educational study."
        },
        {
            "Section": "Project",
            "Metric": "Data type",
            "Result": "Synthetic data",
            "Interpretation": "No real participant records are used."
        },
        {
            "Section": "Edit checks",
            "Metric": "Clean dataset findings",
            "Result": len(clean_findings),
            "Interpretation": "Findings from implemented checks only."
        },
        {
            "Section": "Edit checks",
            "Metric": "Test dataset findings",
            "Result": len(test_findings),
            "Interpretation": "Deliberately seeded defects detected by the edit-check engine."
        },
        {
            "Section": "Edit checks",
            "Metric": "Test findings classified as Error",
            "Result": test_errors,
            "Interpretation": "Requires investigation; not an automatic source-data correction."
        },
        {
            "Section": "Edit checks",
            "Metric": "Test findings classified as Review",
            "Result": test_reviews,
            "Interpretation": "Requires human review."
        },
        {
            "Section": "Queries",
            "Metric": "Generated query drafts",
            "Result": len(queries),
            "Interpretation": "Draft queries derived from eligible test findings."
        },
        {
            "Section": "Query lifecycle",
            "Metric": "Lifecycle demonstration records",
            "Result": len(lifecycle_queries),
            "Interpretation": "Separate working copy used to demonstrate status transitions."
        },
        {
            "Section": "Reconciliation",
            "Metric": "Reconciliation findings",
            "Result": len(reconciliation_findings),
            "Interpretation": "Based on the reconciliation report currently saved."
        },
        {
            "Section": "Report",
            "Metric": "Generated at",
            "Result": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            "Interpretation": "Local computer time."
        },
    ]

    write_table(
        overview,
        ["Section", "Metric", "Result", "Interpretation"],
        overview_rows
    )

    # Test edit-check findings.
    test_sheet = workbook.create_sheet("Test Findings")

    if test_findings:
        test_headers = list(test_findings[0].keys())
        write_table(test_sheet, test_headers, test_findings)
    else:
        write_table(
            test_sheet,
            ["Finding"],
            [{"Finding": "No findings from implemented checks."}]
        )

    # Query drafts.
    query_sheet = workbook.create_sheet("Query Drafts")

    if queries:
        query_headers = list(queries[0].keys())
        write_table(query_sheet, query_headers, queries)
    else:
        write_table(
            query_sheet,
            ["Query"],
            [{"Query": "No query drafts found."}]
        )

    # Query lifecycle demonstration.
    lifecycle_sheet = workbook.create_sheet("Lifecycle Demo")

    if lifecycle_queries:
        lifecycle_headers = list(lifecycle_queries[0].keys())
        write_table(
            lifecycle_sheet,
            lifecycle_headers,
            lifecycle_queries
        )
    else:
        write_table(
            lifecycle_sheet,
            ["Record"],
            [{"Record": "No lifecycle records found."}]
        )

    # Reconciliation findings.
    reconciliation_sheet = workbook.create_sheet("Reconciliation")

    if reconciliation_findings:
        reconciliation_headers = list(
            reconciliation_findings[0].keys()
        )
        write_table(
            reconciliation_sheet,
            reconciliation_headers,
            reconciliation_findings
        )
    else:
        write_table(
            reconciliation_sheet,
            ["Finding"],
            [{"Finding": "No findings in the saved reconciliation report."}]
        )

    # Add a small guide to interpreting the report.
    notes = workbook.create_sheet("Read Me")

    notes_rows = [
        {
            "Topic": "Purpose",
            "Details": "Consolidates selected outputs from this educational CDM workflow."
        },
        {
            "Topic": "Synthetic data",
            "Details": "All study records are fictional and intended for learning."
        },
        {
            "Topic": "Findings",
            "Details": "Findings indicate conditions detected by programmed checks; human review may be needed."
        },
        {
            "Topic": "Query drafts",
            "Details": "Generated queries are training examples and do not represent actual site communication."
        },
        {
            "Topic": "Lifecycle",
            "Details": "The lifecycle workbook demonstrates workflow transitions on a working copy."
        },
        {
            "Topic": "Reconciliation",
            "Details": "Reflects the reconciliation report currently saved. Run reconciliation against the intended dataset before interpreting it."
        },
        {
            "Topic": "Limitations",
            "Details": "This workbook does not establish GCP compliance, production EDC experience, or readiness for a real database lock."
        },
    ]

    write_table(notes, ["Topic", "Details"], notes_rows)

    # Save the consolidated report.
    workbook.save(OUTPUT_FILE)

    print("CDM summary report created successfully.")
    print(f"Output: {OUTPUT_FILE}")
    print()
    print(f"Clean edit-check findings: {len(clean_findings)}")
    print(f"Test edit-check findings: {len(test_findings)}")
    print(f"  Errors: {test_errors}")
    print(f"  Reviews: {test_reviews}")
    print(f"Generated query drafts: {len(queries)}")
    print(f"Lifecycle records: {len(lifecycle_queries)}")
    print(f"Reconciliation findings in saved report: {len(reconciliation_findings)}")
    print()
    print("Note: The summary reflects existing reports; it does not rerun the checks.")


if __name__ == "__main__":
    main()
