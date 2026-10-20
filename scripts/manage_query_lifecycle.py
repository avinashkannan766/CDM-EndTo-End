
from pathlib import Path
from datetime import date
from openpyxl import load_workbook


# --------------------------------------------------
# Project paths
# --------------------------------------------------

BASE_DIR = Path(__file__).resolve().parent.parent
REPORT_DIR = BASE_DIR / "reports"

SOURCE_FILE = REPORT_DIR / "generated_query_log.xlsx"
OUTPUT_FILE = REPORT_DIR / "query_lifecycle_working.xlsx"


# --------------------------------------------------
# Workflow configuration
# --------------------------------------------------

VALID_STATUSES = {
    "Open",
    "Answered",
    "Resolved",
    "Closed",
}

ALLOWED_TRANSITIONS = {
    "Open": {"Answered"},
    "Answered": {"Open", "Resolved"},
    "Resolved": {"Answered", "Closed"},
    "Closed": set(),
}


# --------------------------------------------------
# Helper functions
# --------------------------------------------------

def normalize_text(value):
    if value is None:
        return ""
    return str(value).strip()


def find_column(headers, column_name):
    for index, header in enumerate(headers, start=1):
        if normalize_text(header) == column_name:
            return index

    raise ValueError(
        f"Required column '{column_name}' was not found."
    )


def load_query_workbook():
    if not SOURCE_FILE.exists():
        raise FileNotFoundError(
            f"Source query log not found: {SOURCE_FILE}"
        )

    workbook = load_workbook(SOURCE_FILE)

    if "Generated Queries" not in workbook.sheetnames:
        workbook.close()
        raise ValueError(
            "The source workbook has no 'Generated Queries' sheet."
        )

    sheet = workbook["Generated Queries"]

    if sheet.max_row < 1:
        workbook.close()
        raise ValueError("The query log is empty.")

    headers = [
        normalize_text(
            sheet.cell(row=1, column=col).value
        )
        for col in range(1, sheet.max_column + 1)
    ]

    required_columns = [
        "Query_ID",
        "Query_Status",
        "Response",
        "Resolution",
        "Date_Closed",
        "Reviewer_Notes",
    ]

    column_map = {
        name: find_column(headers, name)
        for name in required_columns
    }

    return workbook, sheet, column_map


# --------------------------------------------------
# Query status transition
# --------------------------------------------------

def transition_query(
    sheet,
    column_map,
    query_id,
    new_status,
    response=None,
    resolution=None,
    reviewer_note=None,
):
    """
    Update one query while enforcing workflow rules.

    All transition validation happens before worksheet cells
    are modified. A rejected transition therefore cannot
    partially change the query in memory.

    This function does not determine whether a clinical data
    issue has actually been corrected. Human review is required.
    """

    query_id = normalize_text(query_id)
    new_status = normalize_text(new_status)

    # Validate the requested status.
    if new_status not in VALID_STATUSES:
        raise ValueError(
            f"Invalid status '{new_status}'. "
            f"Allowed statuses: {sorted(VALID_STATUSES)}"
        )

    query_id_col = column_map["Query_ID"]
    status_col = column_map["Query_Status"]
    response_col = column_map["Response"]
    resolution_col = column_map["Resolution"]
    closed_col = column_map["Date_Closed"]
    notes_col = column_map["Reviewer_Notes"]

    # Locate the query.
    matching_row = None

    for row_number in range(2, sheet.max_row + 1):
        existing_id = normalize_text(
            sheet.cell(
                row=row_number,
                column=query_id_col,
            ).value
        )

        if existing_id == query_id:
            matching_row = row_number
            break

    if matching_row is None:
        raise ValueError(
            f"Query ID '{query_id}' was not found."
        )

    # Read current values without modifying the worksheet.
    status_cell = sheet.cell(
        row=matching_row,
        column=status_col,
    )
    response_cell = sheet.cell(
        row=matching_row,
        column=response_col,
    )
    resolution_cell = sheet.cell(
        row=matching_row,
        column=resolution_col,
    )
    closed_cell = sheet.cell(
        row=matching_row,
        column=closed_col,
    )
    notes_cell = sheet.cell(
        row=matching_row,
        column=notes_col,
    )

    current_status = normalize_text(status_cell.value)

    if current_status not in VALID_STATUSES:
        raise ValueError(
            f"Query '{query_id}' has invalid current status "
            f"'{current_status}'."
        )

    if new_status == current_status:
        raise ValueError(
            f"Query '{query_id}' is already '{current_status}'."
        )

    # Validate the requested transition.
    if new_status not in ALLOWED_TRANSITIONS[current_status]:
        raise ValueError(
            f"Transition '{current_status}' -> '{new_status}' "
            f"is not allowed for query '{query_id}'."
        )

    # Work with proposed values first. Do not write to cells yet.
    proposed_response = (
        response
        if response is not None
        else response_cell.value
    )

    proposed_resolution = (
        resolution
        if resolution is not None
        else resolution_cell.value
    )

    # Validate requirements for the destination status.
    if new_status == "Answered":
        if not normalize_text(proposed_response):
            raise ValueError(
                f"Query '{query_id}' cannot become Answered "
                "without a documented Response."
            )

    if new_status == "Resolved":
        if not normalize_text(proposed_response):
            raise ValueError(
                f"Query '{query_id}' cannot become Resolved "
                "without a documented Response."
            )

        if not normalize_text(proposed_resolution):
            raise ValueError(
                f"Query '{query_id}' cannot become Resolved "
                "without a documented Resolution."
            )

    if new_status == "Closed":
        if current_status != "Resolved":
            raise ValueError(
                f"Query '{query_id}' must be Resolved before closure."
            )

        if not normalize_text(proposed_response):
            raise ValueError(
                f"Query '{query_id}' cannot be Closed without a Response."
            )

        if not normalize_text(proposed_resolution):
            raise ValueError(
                f"Query '{query_id}' cannot be Closed without a Resolution."
            )

    # Prepare the reviewer note before changing any cells.
    proposed_notes = normalize_text(notes_cell.value)

    if reviewer_note:
        if proposed_notes:
            proposed_notes += " | " + normalize_text(reviewer_note)
        else:
            proposed_notes = normalize_text(reviewer_note)

    # All validation has succeeded. Apply changes now.
    response_cell.value = proposed_response
    resolution_cell.value = proposed_resolution
    notes_cell.value = proposed_notes or None

    # Set or clear the closure date.
    if new_status == "Closed":
        closed_cell.value = date.today()
        closed_cell.number_format = "yyyy-mm-dd"

    elif new_status == "Open":
        closed_cell.value = None

    # Update status last.
    status_cell.value = new_status

    print(
        f"{query_id}: {current_status} -> {new_status}"
    )


# --------------------------------------------------
# Workflow demonstration
# --------------------------------------------------

def run_workflow_demo():
    workbook, sheet, column_map = load_query_workbook()

    try:
        # Demonstrate rejection of an invalid transition.
        try:
            transition_query(
                sheet,
                column_map,
                query_id="QRY-001",
                new_status="Closed",
            )
        except ValueError as error:
            print("Expected validation result:")
            print(f"  {error}")

        # Demonstrate the valid lifecycle using fictional text.
        transition_query(
            sheet,
            column_map,
            query_id="QRY-001",
            new_status="Answered",
            response=(
                "TRAINING DEMONSTRATION ONLY: "
                "A hypothetical source review response was entered."
            ),
            reviewer_note=(
                "Lifecycle demonstration; fictional response, "
                "not verified source data."
            ),
        )

        transition_query(
            sheet,
            column_map,
            query_id="QRY-001",
            new_status="Resolved",
            resolution=(
                "TRAINING DEMONSTRATION ONLY: "
                "Hypothetical resolution recorded for workflow testing."
            ),
            reviewer_note=(
                "Resolution shown for software demonstration only."
            ),
        )

        transition_query(
            sheet,
            column_map,
            query_id="QRY-001",
            new_status="Closed",
            reviewer_note=(
                "Closed for lifecycle demonstration only; "
                "not a real clinical record."
            ),
        )

        # Save to the separate working copy.
        workbook.save(OUTPUT_FILE)

        print("\nQuery lifecycle demonstration completed.")
        print(f"Output saved: {OUTPUT_FILE}")
        print(
            "Only the working copy was changed. "
            "The generated query log remains unchanged."
        )

    finally:
        workbook.close()


if __name__ == "__main__":
    run_workflow_demo()
