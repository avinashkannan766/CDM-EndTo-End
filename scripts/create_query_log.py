from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.datavalidation import DataValidation


ROOT = Path(__file__).resolve().parent.parent
OUTPUT_FILE = ROOT / "data" / "query_log.xlsx"
OUTPUT_FILE.parent.mkdir(parents=True, exist_ok=True)

wb = Workbook()

# -------------------------
# Query Log sheet
# -------------------------
ws = wb.active
ws.title = "Query Log"

headers = [
    "Query_ID",
    "Subject_ID",
    "Form_Domain",
    "Record_Reference",
    "Edit_Check_ID",
    "Query_Type",
    "Query_Text",
    "Date_Raised",
    "Query_Status",
    "Response",
    "Resolution",
    "Date_Closed",
    "Reviewer_Notes",
]
ws.append(headers)

# One illustrative example row; not a real clinical query.
ws.append([
    "QRY-EXAMPLE",
    "DM001-004",
    "Visits",
    "DM001-004 / Week 4",
    "EC-009",
    "Missing Data",
    "EXAMPLE ONLY: The visit is marked Completed, but Actual_Date is blank. Please verify the visit date or clarify why it is unavailable.",
    None,
    "Open",
    "",
    "",
    None,
    "Synthetic training example. Replace or remove before using the log for project results.",
])

header_fill = PatternFill("solid", fgColor="1F4E78")
for cell in ws[1]:
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = header_fill
    cell.alignment = Alignment(wrap_text=True, vertical="top")

for row in ws.iter_rows(min_row=2):
    for cell in row:
        cell.alignment = Alignment(wrap_text=True, vertical="top")

widths = {
    "A": 18, "B": 18, "C": 20, "D": 28, "E": 18,
    "F": 20, "G": 70, "H": 16, "I": 18, "J": 55,
    "K": 55, "L": 16, "M": 55,
}
for col, width in widths.items():
    ws.column_dimensions[col].width = width

for row_num in range(2, 202):
    ws.cell(row_num, 8).number_format = "yyyy-mm-dd"
    ws.cell(row_num, 12).number_format = "yyyy-mm-dd"

ws.freeze_panes = "A2"
ws.auto_filter.ref = "A1:M2"
ws.sheet_view.showGridLines = False

table = Table(displayName="QueryLogTable", ref="A1:M2")
table.tableStyleInfo = TableStyleInfo(
    name="TableStyleMedium2",
    showFirstColumn=False,
    showLastColumn=False,
    showRowStripes=True,
    showColumnStripes=False,
)
ws.add_table(table)

# Query status dropdown
status_validation = DataValidation(
    type="list",
    formula1='"Open,Answered,Resolved,Closed"',
    allow_blank=True,
)
status_validation.error = "Choose a status from the dropdown list."
status_validation.errorTitle = "Invalid query status"
status_validation.prompt = "Select the current query status."
status_validation.promptTitle = "Query status"
ws.add_data_validation(status_validation)
status_validation.add("I2:I201")

# Query type dropdown
type_validation = DataValidation(
    type="list",
    formula1='"Missing Data,Inconsistent Data,Invalid Value,Duplicate Record,Date Discrepancy,Other"',
    allow_blank=True,
)
type_validation.error = "Choose a query type from the dropdown list."
type_validation.errorTitle = "Invalid query type"
ws.add_data_validation(type_validation)
type_validation.add("F2:F201")

# -------------------------
# Workflow Guide sheet
# -------------------------
guide = wb.create_sheet("Workflow Guide")
guide_rows = [
    ["Stage", "Meaning", "Expected action"],
    ["Open", "A finding has been documented and clarification is requested.", "Record the query text, related check, participant, domain, and date raised."],
    ["Answered", "A response has been recorded.", "Review whether the response resolves the finding; request clarification if needed."],
    ["Resolved", "The response and any required data correction have been assessed.", "Record the resolution and ensure any correction is traceable."],
    ["Closed", "The reviewer confirms that no further action is required.", "Record the closure date and reviewer notes."],
    ["Important", "A query should not be closed solely because a response was received.", "Document the rationale and preserve the original raw data."],
    ["Synthetic example", "The example row in Query Log is instructional only.", "Do not count it as a genuine detected defect or a completed project query."],
]
for row in guide_rows:
    guide.append(row)

for cell in guide[1]:
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = header_fill
    cell.alignment = Alignment(wrap_text=True, vertical="top")

for row in guide.iter_rows(min_row=2):
    for cell in row:
        cell.alignment = Alignment(wrap_text=True, vertical="top")

guide.column_dimensions["A"].width = 20
guide.column_dimensions["B"].width = 55
guide.column_dimensions["C"].width = 90
guide.freeze_panes = "A2"
guide.sheet_view.showGridLines = False

guide_table = Table(displayName="WorkflowGuideTable", ref=guide.dimensions)
guide_table.tableStyleInfo = TableStyleInfo(
    name="TableStyleMedium2",
    showFirstColumn=False,
    showLastColumn=False,
    showRowStripes=True,
    showColumnStripes=False,
)
guide.add_table(guide_table)

wb.save(OUTPUT_FILE)

print(f"Created query log: {OUTPUT_FILE}")
print("Sheets: Query Log, Workflow Guide")
print("Query status dropdown: Open, Answered, Resolved, Closed")
print("Query type dropdown configured.")
print("One clearly labelled synthetic example row included.")
