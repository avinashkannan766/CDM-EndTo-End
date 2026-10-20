
from pathlib import Path
from collections import Counter, defaultdict
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, Alignment


# --------------------------------------------------
# Project paths and study configuration
# --------------------------------------------------

BASE_DIR = Path(__file__).resolve().parent.parent
DATA_DIR = BASE_DIR / "data"
REPORT_DIR = BASE_DIR / "reports"

INPUT_FILE = DATA_DIR / "clinical_study_clean.xlsx"
OUTPUT_FILE = REPORT_DIR / "reconciliation_report.xlsx"

EXPECTED_VISITS = [
    "Screening",
    "Baseline",
    "Week 4",
    "Week 8",
    "Week 12",
]

FORM_STATUS_COLUMNS = [
    "Visit_Form_Status",
    "Vitals_Form_Status",
    "Laboratory_Form_Status",
    "AE_Form_Status",
    "Medication_Form_Status",
]

ALLOWED_FORM_STATUSES = {
    "Complete",
    "Incomplete",
    "Not Started",
    "Not Applicable",
}

REPORT_DIR.mkdir(parents=True, exist_ok=True)


# --------------------------------------------------
# Workbook reading helpers
# --------------------------------------------------

def read_sheet(workbook, sheet_name):
    if sheet_name not in workbook.sheetnames:
        raise ValueError(
            f"Required sheet '{sheet_name}' not found. "
            f"Available sheets: {workbook.sheetnames}"
        )

    sheet = workbook[sheet_name]
    rows = list(sheet.iter_rows(values_only=True))

    if not rows:
        return []

    headers = [
        str(value).strip() if value is not None else ""
        for value in rows[0]
    ]

    records = []

    for values in rows[1:]:
        if all(value is None for value in values):
            continue

        records.append({
            headers[index]: values[index]
            for index in range(min(len(headers), len(values)))
            if headers[index]
        })

    return records


def normalize(value):
    if value is None:
        return ""
    return str(value).strip()


# --------------------------------------------------
# Reconciliation engine
# --------------------------------------------------

def run_reconciliation(input_file):
    if not input_file.exists():
        raise FileNotFoundError(
            f"Input workbook not found: {input_file}"
        )

    workbook = load_workbook(input_file, data_only=True)

    try:
        demographics = read_sheet(workbook, "Demographics")
        visits = read_sheet(workbook, "Visits")
        form_tracking = read_sheet(workbook, "Form_Tracking")

    finally:
        workbook.close()

    findings = []

    # Establish the expected participant population.
    subjects = sorted({
        normalize(row.get("Subject_ID"))
        for row in demographics
        if normalize(row.get("Subject_ID"))
    })

    if not subjects:
        raise ValueError(
            "No participant IDs were found in Demographics."
        )

    expected_pairs = {
        (subject_id, visit_name)
        for subject_id in subjects
        for visit_name in EXPECTED_VISITS
    }

    # Index actual visit records.
    visit_groups = defaultdict(list)

    for row in visits:
        subject_id = normalize(row.get("Subject_ID"))
        visit_name = normalize(row.get("Visit_Name"))

        if subject_id and visit_name:
            visit_groups[(subject_id, visit_name)].append(row)

    actual_pairs = set(visit_groups)

    # Index form tracking rows.
    tracking_groups = defaultdict(list)

    for row in form_tracking:
        subject_id = normalize(row.get("Subject_ID"))
        visit_name = normalize(row.get("Visit_Name"))

        if subject_id and visit_name:
            tracking_groups[(subject_id, visit_name)].append(row)

    actual_tracking_pairs = set(tracking_groups)

    def add_finding(
        finding_type,
        subject_id,
        visit_name,
        record_reference,
        description,
        severity="Review",
    ):
        findings.append({
            "Finding_Type": finding_type,
            "Subject_ID": subject_id,
            "Visit_Name": visit_name,
            "Record_Reference": record_reference,
            "Severity": severity,
            "Finding_Description": description,
        })

    # 1. Expected visit records missing from the Visits sheet.
    for subject_id, visit_name in sorted(
        expected_pairs - actual_pairs
    ):
        add_finding(
            "Missing Visit Record",
            subject_id,
            visit_name,
            f"{subject_id} / {visit_name}",
            "An expected participant-visit combination is absent "
            "from the Visits sheet."
        )

    # 2. Visit records outside the configured expected schedule.
    for subject_id, visit_name in sorted(
        actual_pairs - expected_pairs
    ):
        add_finding(
            "Unexpected Visit Record",
            subject_id,
            visit_name,
            f"{subject_id} / {visit_name}",
            "A participant-visit combination exists outside the "
            "configured expected visit schedule."
        )

    # 3. Duplicate participant-visit records.
    for (subject_id, visit_name), rows in sorted(
        visit_groups.items()
    ):
        if len(rows) > 1:
            add_finding(
                "Duplicate Visit Record",
                subject_id,
                visit_name,
                f"{subject_id} / {visit_name}",
                f"Found {len(rows)} rows for the same participant "
                "and visit combination."
            )

    # 4. Expected tracking rows missing.
    for subject_id, visit_name in sorted(
        expected_pairs - actual_tracking_pairs
    ):
        add_finding(
            "Missing Form Tracking Record",
            subject_id,
            visit_name,
            f"{subject_id} / {visit_name}",
            "No Form_Tracking row exists for an expected "
            "participant-visit combination."
        )

    # 5. Duplicate tracking rows.
    for (subject_id, visit_name), rows in sorted(
        tracking_groups.items()
    ):
        if len(rows) > 1:
            add_finding(
                "Duplicate Form Tracking Record",
                subject_id,
                visit_name,
                f"{subject_id} / {visit_name}",
                f"Found {len(rows)} tracking rows for the same "
                "participant and visit combination."
            )

    # 6. Check form statuses in every tracking row.
    for (subject_id, visit_name), rows in sorted(
        tracking_groups.items()
    ):
        for row_index, row in enumerate(rows, start=1):
            record_reference = (
                f"{subject_id} / {visit_name} / "
                f"Tracking row {row_index}"
            )

            for column_name in FORM_STATUS_COLUMNS:
                status = normalize(row.get(column_name))

                if not status:
                    add_finding(
                        "Missing Form Status",
                        subject_id,
                        visit_name,
                        record_reference,
                        f"{column_name} is blank."
                    )
                elif status not in ALLOWED_FORM_STATUSES:
                    add_finding(
                        "Invalid Form Status",
                        subject_id,
                        visit_name,
                        record_reference,
                        f"{column_name} contains unsupported value "
                        f"'{status}'."
                    )

            review_status = normalize(row.get("Review_Status"))

            if not review_status:
                add_finding(
                    "Missing Review Status",
                    subject_id,
                    visit_name,
                    record_reference,
                    "Review_Status is blank."
                )
            elif review_status not in {"Reviewed", "Pending Review"}:
                add_finding(
                    "Invalid Review Status",
                    subject_id,
                    visit_name,
                    record_reference,
                    f"Review_Status contains unsupported value "
                    f"'{review_status}'."
                )

    # 7. Flag a tracking row marked complete when its visit
    # record is absent or not marked Completed.
    for (subject_id, visit_name), rows in sorted(
        tracking_groups.items()
    ):
        matching_visits = visit_groups.get(
            (subject_id, visit_name), []
        )

        for row_index, tracking_row in enumerate(rows, start=1):
            visit_form_status = normalize(
                tracking_row.get("Visit_Form_Status")
            ).lower()

            if visit_form_status != "complete":
                continue

            if not matching_visits:
                add_finding(
                    "Completion Without Visit Record",
                    subject_id,
                    visit_name,
                    f"{subject_id} / {visit_name}",
                    "Visit_Form_Status is Complete, but no matching "
                    "visit record exists."
                )
                continue

            if not any(
                normalize(row.get("Visit_Status")).lower()
                == "completed"
                for row in matching_visits
            ):
                add_finding(
                    "Completion Status Mismatch",
                    subject_id,
                    visit_name,
                    f"{subject_id} / {visit_name}",
                    "Visit_Form_Status is Complete, but no matching "
                    "Visits record is marked Completed."
                )

    # Deterministic IDs and ordering.
    findings.sort(
        key=lambda item: (
            item["Finding_Type"],
            item["Subject_ID"],
            item["Visit_Name"],
            item["Record_Reference"],
        )
    )

    for index, finding in enumerate(findings, start=1):
        finding["Finding_ID"] = f"REC-{index:03d}"

    summary = {
        "Participants": len(subjects),
        "Expected Visit Combinations": len(expected_pairs),
        "Visit Rows": sum(len(rows) for rows in visit_groups.values()),
        "Unique Visit Combinations": len(actual_pairs),
        "Form Tracking Rows": sum(
            len(rows) for rows in tracking_groups.values()
        ),
        "Unique Tracking Combinations": len(actual_tracking_pairs),
        "Reconciliation Findings": len(findings),
    }

    return summary, findings


# --------------------------------------------------
# Excel report generation
# --------------------------------------------------

def write_report(summary, findings, output_file):
    workbook = Workbook()

    summary_sheet = workbook.active
    summary_sheet.title = "Summary"
    summary_sheet.append(["Reconciliation Metric", "Value"])

    for metric, value in summary.items():
        summary_sheet.append([metric, value])

    summary_sheet.append([])
    summary_sheet.append([
        "Scope Note",
        "This report checks synthetic participant-visit coverage, "
        "tracking rows, selected status values, and visit completion "
        "consistency. It does not establish clinical database lock "
        "readiness or verify source documents."
    ])

    detail_sheet = workbook.create_sheet("Findings")

    columns = [
        "Finding_ID",
        "Finding_Type",
        "Subject_ID",
        "Visit_Name",
        "Record_Reference",
        "Severity",
        "Finding_Description",
    ]

    detail_sheet.append(columns)

    for finding in findings:
        detail_sheet.append([
            finding.get(column)
            for column in columns
        ])

    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions

        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True
            )

        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(
                    vertical="top",
                    wrap_text=True
                )

        for column_cells in sheet.columns:
            max_length = max(
                len(str(cell.value or ""))
                for cell in column_cells
            )
            column_letter = column_cells[0].column_letter
            sheet.column_dimensions[column_letter].width = min(
                max(max_length + 2, 14),
                70
            )

    workbook.save(output_file)
    workbook.close()


# --------------------------------------------------
# Main
# --------------------------------------------------

if __name__ == "__main__":
    summary, findings = run_reconciliation(INPUT_FILE)
    write_report(summary, findings, OUTPUT_FILE)

    print("\nReconciliation completed.")
    for metric, value in summary.items():
        print(f"{metric}: {value}")

    print(f"\nReport saved: {OUTPUT_FILE}")

    if findings:
        print(
            "\nReview the Findings sheet. "
            "Findings are prompts for review, not automatic "
            "determinations that source data are wrong."
        )
    else:
        print(
            "\nNo findings were generated by the implemented checks."
        )
