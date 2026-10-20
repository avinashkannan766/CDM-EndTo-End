from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo


PROJECT_ROOT = Path(__file__).resolve().parent.parent
OUTPUT_FILE = PROJECT_ROOT / "docs" / "data_dictionary.xlsx"


# Columns:
# Form ID, Field Name, Label, Definition, Data Type, Required,
# Allowed Values, Validation Rule, Notes
fields = [
    # Demographics and consent
    ("DM001-DMOGRAPHICS", "Subject_ID", "Subject ID", "Unique fictional participant identifier.", "Text", "Yes", "Unique nonblank ID", "Must be unique and nonblank.", ""),
    ("DM001-DMOGRAPHICS", "Age", "Age", "Participant age in years.", "Integer", "Yes", "18-70", "Must be an integer from 18 to 70.", "Training assumption only."),
    ("DM001-DMOGRAPHICS", "Sex", "Sex", "Participant sex code as recorded in the simulation.", "Coded text", "Yes", "Female; Male; Other; Unknown", "Must match the documented code list.", "Synthetic data only."),
    ("DM001-DMOGRAPHICS", "Consent_Status", "Consent status", "Whether consent is documented in the simulation.", "Coded text", "Yes", "Yes; No", "Must be Yes or No.", ""),
    ("DM001-DMOGRAPHICS", "Consent_Date", "Consent date", "Date consent was recorded.", "Date", "Conditional", "YYYY-MM-DD", "Required when Consent_Status is Yes.", "Do not invent missing dates."),

    # Eligibility
    ("DM001-ELIGIBILITY", "Subject_ID", "Subject ID", "Participant linked to the eligibility assessment.", "Text", "Yes", "Existing Subject_ID", "Must match the demographics form.", ""),
    ("DM001-ELIGIBILITY", "Visit_Name", "Visit name", "Visit at which eligibility is assessed.", "Coded text", "Yes", "Screening", "Must be Screening.", ""),
    ("DM001-ELIGIBILITY", "Inclusion_1", "Inclusion criterion 1", "Simulated response to the first inclusion criterion.", "Coded text", "Yes", "Yes; No", "Must be Yes or No.", "Simplified training criterion."),
    ("DM001-ELIGIBILITY", "Inclusion_2", "Inclusion criterion 2", "Simulated response to the second inclusion criterion.", "Coded text", "Yes", "Yes; No", "Must be Yes or No.", "Simplified training criterion."),
    ("DM001-ELIGIBILITY", "Exclusion_1", "Exclusion criterion 1", "Simulated response to the exclusion criterion.", "Coded text", "Yes", "Yes; No", "Must be Yes or No.", "Simplified training criterion."),
    ("DM001-ELIGIBILITY", "Eligibility_Status", "Eligibility status", "Overall simulated eligibility outcome.", "Coded text", "Yes", "Eligible; Not Eligible; Pending", "Must match the documented code list.", ""),
    ("DM001-ELIGIBILITY", "Assessment_Date", "Assessment date", "Date of eligibility assessment.", "Date", "Yes", "YYYY-MM-DD", "Must be a valid date.", ""),

    # Visit tracking
    ("DM001-VISITS", "Subject_ID", "Subject ID", "Participant associated with the visit.", "Text", "Yes", "Existing Subject_ID", "Must match the demographics form.", ""),
    ("DM001-VISITS", "Visit_Name", "Visit name", "Scheduled study visit.", "Coded text", "Yes", "Screening; Baseline; Week 4; Week 8; Week 12", "Must match the visit schedule.", ""),
    ("DM001-VISITS", "Planned_Date", "Planned date", "Planned date of the visit.", "Date", "Yes", "YYYY-MM-DD", "Must be a valid date.", ""),
    ("DM001-VISITS", "Actual_Date", "Actual date", "Date the visit actually occurred.", "Date", "Conditional", "YYYY-MM-DD", "Required when Visit_Status is Completed.", ""),
    ("DM001-VISITS", "Visit_Status", "Visit status", "Completion status of the visit.", "Coded text", "Yes", "Not Started; Completed; Missed; Not Done", "Must match the documented code list.", ""),
    ("DM001-VISITS", "Form_Status", "Form status", "Completion/review status of the visit form.", "Coded text", "Yes", "Incomplete; Complete; Reviewed", "Must match the documented code list.", "Form review is separate from visit completion."),

    # Vital signs
    ("DM001-VITALS", "Subject_ID", "Subject ID", "Participant associated with the measurement.", "Text", "Yes", "Existing Subject_ID", "Must match the demographics form.", ""),
    ("DM001-VITALS", "Visit_Name", "Visit name", "Visit at which vital signs were collected.", "Coded text", "Yes", "Screening; Baseline; Week 4; Week 8; Week 12", "Must match the visit schedule.", ""),
    ("DM001-VITALS", "Measurement_Date", "Measurement date", "Date vital signs were measured.", "Date", "Yes", "YYYY-MM-DD", "Must be a valid date.", ""),
    ("DM001-VITALS", "Systolic_BP", "Systolic blood pressure", "Systolic blood pressure value.", "Numeric", "Yes", "Positive numeric value", "Must be numeric and greater than zero; unusual values flagged for review.", "Training rule, not a clinical threshold."),
    ("DM001-VITALS", "Diastolic_BP", "Diastolic blood pressure", "Diastolic blood pressure value.", "Numeric", "Yes", "Positive numeric value", "Must be numeric and greater than zero; unusual values flagged for review.", "Training rule, not a clinical threshold."),
    ("DM001-VITALS", "Heart_Rate", "Heart rate", "Heart rate value.", "Numeric", "Yes", "Positive numeric value", "Must be numeric and greater than zero; unusual values flagged for review.", "Training rule, not a clinical threshold."),
    ("DM001-VITALS", "Weight_kg", "Weight (kg)", "Weight in kilograms.", "Numeric", "Yes", "Positive numeric value", "Must be numeric and greater than zero.", ""),

    # Laboratory results
    ("DM001-LABS", "Lab_ID", "Laboratory record ID", "Unique identifier for a laboratory result record.", "Text", "Yes", "Unique nonblank ID", "Must be unique and nonblank.", ""),
    ("DM001-LABS", "Subject_ID", "Subject ID", "Participant associated with the laboratory result.", "Text", "Yes", "Existing Subject_ID", "Must match the demographics form.", ""),
    ("DM001-LABS", "Visit_Name", "Visit name", "Visit associated with the laboratory collection.", "Coded text", "Yes", "Screening; Baseline; Week 4; Week 8; Week 12", "Must match the visit schedule.", ""),
    ("DM001-LABS", "Test_Name", "Test name", "Name of the laboratory test.", "Coded text", "Yes", "Glucose; HbA1c; Creatinine", "Must match the finalized study test list.", "Illustrative test list."),
    ("DM001-LABS", "Result_Value", "Result value", "Numeric laboratory result.", "Numeric", "Yes", "Numeric", "Must be numeric when a result is recorded.", ""),
    ("DM001-LABS", "Unit", "Unit", "Unit associated with the result.", "Text", "Yes", "Test-specific unit", "Must match the unit defined for the test.", ""),
    ("DM001-LABS", "Collection_Date", "Collection date", "Date the specimen was collected.", "Date", "Yes", "YYYY-MM-DD", "Must be a valid date.", ""),

    # Adverse events
    ("DM001-AE", "AE_ID", "Adverse event ID", "Unique identifier for an adverse event record.", "Text", "Yes", "Unique nonblank ID", "Must be unique and nonblank.", ""),
    ("DM001-AE", "Subject_ID", "Subject ID", "Participant associated with the event.", "Text", "Yes", "Existing Subject_ID", "Must match the demographics form.", ""),
    ("DM001-AE", "Event_Description", "Event description", "Description of the simulated adverse event.", "Text", "Yes", "Nonblank text", "Must not be blank.", ""),
    ("DM001-AE", "Onset_Date", "Onset date", "Date the event began.", "Date", "Yes", "YYYY-MM-DD", "Must be a valid date.", ""),
    ("DM001-AE", "Seriousness", "Seriousness", "Whether the event is classified as serious in the simulation.", "Coded text", "Yes", "Yes; No", "Must be Yes or No.", "Not a real safety assessment."),
    ("DM001-AE", "Outcome", "Outcome", "Recorded outcome of the event.", "Coded text", "Yes", "Resolved; Recovering; Ongoing; Unknown", "Must match the documented code list.", ""),

    # Concomitant medications
    ("DM001-MEDICATIONS", "Medication_ID", "Medication record ID", "Unique identifier for a medication record.", "Text", "Yes", "Unique nonblank ID", "Must be unique and nonblank.", ""),
    ("DM001-MEDICATIONS", "Subject_ID", "Subject ID", "Participant associated with the medication.", "Text", "Yes", "Existing Subject_ID", "Must match the demographics form.", ""),
    ("DM001-MEDICATIONS", "Medication_Name", "Medication name", "Name of the simulated concomitant medication.", "Text", "Yes", "Nonblank text", "Must not be blank.", ""),
    ("DM001-MEDICATIONS", "Start_Date", "Start date", "Date medication use began.", "Date", "Yes", "YYYY-MM-DD", "Must be a valid date.", ""),
    ("DM001-MEDICATIONS", "End_Date", "End date", "Date medication use ended.", "Date", "Conditional", "YYYY-MM-DD or blank if ongoing", "If present, must be a valid date and not precede Start_Date.", ""),
    ("DM001-MEDICATIONS", "Indication", "Indication", "Reason the medication was used in the simulation.", "Text", "Yes", "Nonblank text", "Must not be blank.", ""),
]

code_lists = [
    ("Sex", "Female", "Synthetic training code list"),
    ("Sex", "Male", "Synthetic training code list"),
    ("Sex", "Other", "Synthetic training code list"),
    ("Sex", "Unknown", "Synthetic training code list"),
    ("Consent_Status", "Yes", "Consent documented"),
    ("Consent_Status", "No", "Consent not documented"),
    ("Yes_No", "Yes", "Generic binary response"),
    ("Yes_No", "No", "Generic binary response"),
    ("Eligibility_Status", "Eligible", "Training status"),
    ("Eligibility_Status", "Not Eligible", "Training status"),
    ("Eligibility_Status", "Pending", "Training status"),
    ("Visit_Name", "Screening", "Scheduled visit"),
    ("Visit_Name", "Baseline", "Scheduled visit"),
    ("Visit_Name", "Week 4", "Scheduled visit"),
    ("Visit_Name", "Week 8", "Scheduled visit"),
    ("Visit_Name", "Week 12", "Scheduled visit"),
    ("Visit_Status", "Not Started", "Visit has not started"),
    ("Visit_Status", "Completed", "Visit completed"),
    ("Visit_Status", "Missed", "Visit missed"),
    ("Visit_Status", "Not Done", "Visit not performed"),
    ("Form_Status", "Incomplete", "Form not complete"),
    ("Form_Status", "Complete", "Form completion recorded"),
    ("Form_Status", "Reviewed", "Form reviewed"),
    ("Test_Name", "Glucose", "Illustrative lab test"),
    ("Test_Name", "HbA1c", "Illustrative lab test"),
    ("Test_Name", "Creatinine", "Illustrative lab test"),
    ("Seriousness", "Yes", "Training code"),
    ("Seriousness", "No", "Training code"),
    ("Outcome", "Resolved", "Training code"),
    ("Outcome", "Recovering", "Training code"),
    ("Outcome", "Ongoing", "Training code"),
    ("Outcome", "Unknown", "Training code"),
]

wb = Workbook()
ws = wb.active
ws.title = "Data Dictionary"

headers = [
    "Form_ID", "Field_Name", "Field_Label", "Definition", "Data_Type",
    "Required_Status", "Allowed_Values", "Validation_Rule", "Notes"
]
ws.append(headers)

for row in fields:
    ws.append(row)

# Style the dictionary
header_fill = PatternFill("solid", fgColor="1F4E78")
for cell in ws[1]:
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = header_fill
    cell.alignment = Alignment(wrap_text=True, vertical="top")

for row in ws.iter_rows(min_row=2):
    for cell in row:
        cell.alignment = Alignment(wrap_text=True, vertical="top")

widths = [24, 23, 28, 45, 16, 18, 38, 48, 34]
for idx, width in enumerate(widths, start=1):
    ws.column_dimensions[ws.cell(1, idx).column_letter].width = width

ws.freeze_panes = "A2"
ws.auto_filter.ref = ws.dimensions
ws.sheet_view.showGridLines = False

table = Table(displayName="DataDictionaryTable", ref=ws.dimensions)
table.tableStyleInfo = TableStyleInfo(
    name="TableStyleMedium2",
    showFirstColumn=False,
    showLastColumn=False,
    showRowStripes=True,
    showColumnStripes=False,
)
ws.add_table(table)

# Add controlled terminology / code lists
codes_ws = wb.create_sheet("Code Lists")
codes_ws.append(["Code_List_Name", "Allowed_Value", "Description"])
for row in code_lists:
    codes_ws.append(row)

for cell in codes_ws[1]:
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = header_fill
    cell.alignment = Alignment(wrap_text=True, vertical="top")

for row in codes_ws.iter_rows(min_row=2):
    for cell in row:
        cell.alignment = Alignment(wrap_text=True, vertical="top")

for col, width in {"A": 25, "B": 25, "C": 40}.items():
    codes_ws.column_dimensions[col].width = width

codes_ws.freeze_panes = "A2"
codes_ws.auto_filter.ref = codes_ws.dimensions
codes_ws.sheet_view.showGridLines = False

codes_table = Table(displayName="CodeListsTable", ref=codes_ws.dimensions)
codes_table.tableStyleInfo = TableStyleInfo(
    name="TableStyleMedium2",
    showFirstColumn=False,
    showLastColumn=False,
    showRowStripes=True,
    showColumnStripes=False,
)
codes_ws.add_table(codes_table)

OUTPUT_FILE.parent.mkdir(parents=True, exist_ok=True)
wb.save(OUTPUT_FILE)

print(f"Created data dictionary: {OUTPUT_FILE}")
print(f"Data dictionary fields: {len(fields)}")
print(f"Code-list entries: {len(code_lists)}")
