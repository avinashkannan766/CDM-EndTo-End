
from pathlib import Path
from datetime import date
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


# --------------------------------------------------
# 1. File paths
# --------------------------------------------------

BASE_DIR = Path(__file__).resolve().parent.parent
REPORT_DIR = BASE_DIR / "reports"

QUERY_FILE = REPORT_DIR / "query_lifecycle_working.xlsx"
OUTPUT_FILE = REPORT_DIR / "data_correction_audit_log.xlsx"


# --------------------------------------------------
# 2. Workbook column definitions
# --------------------------------------------------

AUDIT_COLUMNS = [
    "Change_ID",
    "Query_ID",
    "Subject_ID",
    "Domain",
    "Record_ID",
    "Field_Name",
    "Old_Value",
    "Proposed_New_Value",
    "Correction_Reason",
    "Review_Status",
    "Reviewed_By",
    "Review_Date",
    "Change_Date",
    "Reviewer_Comments",
]

QUERY_REFERENCE_COLUMNS = [
    "Query_ID",
    "Subject_ID",
    "Form_Domain",
    "Record_Reference",
    "Query_Status",
    "Audit_Record_Required",
]


# --------------------------------------------------
# 3. Formatting helpers
# --------------------------------------------------

def format_worksheet(sheet):
    """Apply basic readable formatting to a worksheet."""

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions

    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(
            fill_type="solid",
            fgColor="1F4E78",
        )
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )

    for column_cells in sheet.columns:
        column_letter = get_column_letter(column_cells[0].column)
        max_length = 0

        for cell in column_cells:
            if cell.value is not None:
                max_length = max(
                    max_length,
                    len(str(cell.value)),
                )

        sheet.column_dimensions[column_letter].width = min(
            max(max_length + 2, 14),
            40,
        )

    sheet.row_dimensions[1].height = 32


def add_headers(sheet, columns):
    """Add and format a worksheet's header row."""

    sheet.append(columns)
    format_worksheet(sheet)


# --------------------------------------------------
# 4. Load query references
# --------------------------------------------------

def load_query_references():
    """
    Read query information for traceability.

    The source workbook is opened read-only.
    No source data or query statuses are changed.
    """

    if not QUERY_FILE.exists():
        raise FileNotFoundError(
            f"Query lifecycle workbook not found:\n{QUERY_FILE}\n"
            "Run scripts/manage_query_lifecycle.py first."
        )

    workbook = load_workbook(
        QUERY_FILE,
        read_only=True,
        data_only=True,
    )

    try:
        if "Generated Queries" not in workbook.sheetnames:
            raise ValueError(
                "Expected worksheet 'Generated Queries' "
                "was not found."
            )

        sheet = workbook["Generated Queries"]

        headers = [
            str(cell.value).strip()
            if cell.value is not None
            else ""
            for cell in sheet[1]
        ]

        required_columns = [
            "Query_ID",
            "Subject_ID",
            "Form_Domain",
            "Record_Reference",
            "Query_Status",
        ]

        missing = [
            column
            for column in required_columns
            if column not in headers
        ]

        if missing:
            raise ValueError(
                "Missing required query-log columns: "
                + ", ".join(missing)
            )

        positions = {
            column: headers.index(column)
            for column in required_columns
        }

        references = []

        for row in sheet.iter_rows(
            min_row=2,
            values_only=True,
        ):
            query_id = row[positions["Query_ID"]]

            if query_id is None or not str(query_id).strip():
                continue

            references.append([
                str(query_id).strip(),
                row[positions["Subject_ID"]],
                row[positions["Form_Domain"]],
                row[positions["Record_Reference"]],
                row[positions["Query_Status"]],
                "Pending review",
            ])

        return references

    finally:
        workbook.close()


# --------------------------------------------------
# 5. Create the correction audit workbook
# --------------------------------------------------

def create_audit_workbook():
    """Create a separate correction audit workbook."""

    if OUTPUT_FILE.exists():
        raise FileExistsError(
            f"Output workbook already exists:\n{OUTPUT_FILE}\n"
            "It has NOT been overwritten. Rename or move the "
            "existing file if you intentionally want to "
            "generate a new version."
        )

    query_references = load_query_references()

    workbook = Workbook()

    # Worksheet 1: Correction audit trail
    audit_sheet = workbook.active
    audit_sheet.title = "Correction Audit"

    add_headers(audit_sheet, AUDIT_COLUMNS)

    # No real correction is being claimed here.
    # This sheet is a blank template for proposed changes.
    audit_sheet.append([
        None,   # Change_ID
        None,   # Query_ID
        None,   # Subject_ID
        None,   # Domain
        None,   # Record_ID
        None,   # Field_Name
        None,   # Old_Value
        None,   # Proposed_New_Value
        None,   # Correction_Reason
        "Pending Review",
        None,   # Reviewed_By
        None,   # Review_Date
        None,   # Change_Date
        "No correction has been applied.",
    ])

    # Worksheet 2: Query traceability
    reference_sheet = workbook.create_sheet(
        "Query Reference"
    )

    add_headers(
        reference_sheet,
        QUERY_REFERENCE_COLUMNS,
    )

    for reference in query_references:
        reference_sheet.append(reference)

    # Worksheet 3: Instructions
    instructions_sheet = workbook.create_sheet(
        "Instructions"
    )

    instructions = [
        ["Data Correction Audit Trail - Instructions"],
        [
            "Purpose",
            "Document proposed and reviewed data corrections "
            "with traceability to a query.",
        ],
        [
            "Source",
            "reports/query_lifecycle_working.xlsx",
        ],
        [
            "Data integrity",
            "The source query log and clinical data workbooks "
            "are not modified by this script.",
        ],
        [
            "Important",
            "Closing a query does not prove that the underlying "
            "data has been corrected.",
        ],
        [
            "Change_ID",
            "Unique identifier for each proposed correction.",
        ],
        [
            "Query_ID",
            "Identifier of the related data query, where applicable.",
        ],
        [
            "Old_Value",
            "Original value recorded before an approved correction.",
        ],
        [
            "Proposed_New_Value",
            "Suggested replacement value; do not guess missing data.",
        ],
        [
            "Correction_Reason",
            "Documented reason and supporting evidence for the change.",
        ],
        [
            "Review_Status",
            "Use Pending Review, Approved, or Rejected.",
        ],
        [
            "Reviewed_By",
            "Name or role of the authorized reviewer.",
        ],
        [
            "Review_Date",
            "Date on which the proposed correction was reviewed.",
        ],
        [
            "Change_Date",
            "Date the approved correction was actually applied, "
            "if applicable.",
        ],
        [
            "Training limitation",
            "This is an educational template using synthetic "
            "study data. It does not represent a production "
            "clinical database or a real correction.",
        ],
    ]

    for instruction in instructions:
        instructions_sheet.append(instruction)

    instructions_sheet.column_dimensions["A"].width = 25
    instructions_sheet.column_dimensions["B"].width = 90

    for row in instructions_sheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True,
            )

    instructions_sheet["A1"].font = Font(
        bold=True,
        size=14,
        color="FFFFFF",
    )
    instructions_sheet["A1"].fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78",
    )

    for cell in instructions_sheet[1]:
        cell.fill = PatternFill(
            fill_type="solid",
            fgColor="1F4E78",
        )
        cell.font = Font(
            bold=True,
            color="FFFFFF",
        )

    for sheet in workbook.worksheets:
        if sheet.title != "Instructions":
            format_worksheet(sheet)

    REPORT_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    workbook.save(OUTPUT_FILE)
    workbook.close()

    print("Correction audit workbook created successfully.")
    print(f"Output: {OUTPUT_FILE}")
    print(f"Query references loaded: {len(query_references)}")
    print("Source query workbook: unchanged")
    print("Clinical data workbooks: unchanged")
    print("No actual data corrections were applied.")


# --------------------------------------------------
# 6. Run
# --------------------------------------------------

if __name__ == "__main__":
    try:
        create_audit_workbook()

    except (FileNotFoundError, FileExistsError, ValueError) as error:
        print(f"ERROR: {error}")
        raise SystemExit(1)
