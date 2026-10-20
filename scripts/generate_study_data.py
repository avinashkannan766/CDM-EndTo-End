from pathlib import Path
from copy import deepcopy
from datetime import date, timedelta
import random

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo


ROOT = Path(__file__).resolve().parent.parent
DATA_DIR = ROOT / "data"
DOCS_DIR = ROOT / "docs"

DATA_DIR.mkdir(parents=True, exist_ok=True)
DOCS_DIR.mkdir(parents=True, exist_ok=True)

random.seed(2026)

SUBJECT_COUNT = 30
VISIT_NAMES = ["Screening", "Baseline", "Week 4", "Week 8", "Week 12"]
VISIT_OFFSETS = [0, 7, 35, 63, 91]

# Fixed seed and fixed dates make the output reproducible.
start_date = date(2026, 1, 5)

subjects = [f"DM001-{i:03d}" for i in range(1, SUBJECT_COUNT + 1)]
arms = ["Drug X"] * 15 + ["Placebo"] * 15
random.shuffle(arms)

demographics = []
eligibility = []
visits = []
vitals = []
laboratory = []
adverse_events = []
medications = []
form_tracking = []

# -------------------------
# Demographics and consent
# -------------------------
for i, subject_id in enumerate(subjects):
    screening_date = start_date + timedelta(days=i)
    demographics.append({
        "Subject_ID": subject_id,
        "Age": random.randint(30, 68),
        "Sex": random.choice(["Female", "Male"]),
        "Consent_Status": "Yes",
        "Consent_Date": screening_date,
        "Treatment_Arm": arms[i],
    })

    eligibility.append({
        "Subject_ID": subject_id,
        "Visit_Name": "Screening",
        "Inclusion_1": "Yes",
        "Inclusion_2": "Yes",
        "Exclusion_1": "No",
        "Eligibility_Status": "Eligible",
        "Assessment_Date": screening_date,
    })

# -------------------------
# Visits, vital signs, labs
# -------------------------
lab_counter = 1

for i, subject_id in enumerate(subjects):
    screening_date = start_date + timedelta(days=i)

    for visit_name, offset in zip(VISIT_NAMES, VISIT_OFFSETS):
        planned_date = screening_date + timedelta(days=offset)

        visits.append({
            "Subject_ID": subject_id,
            "Visit_Name": visit_name,
            "Planned_Date": planned_date,
            "Actual_Date": planned_date,
            "Visit_Status": "Completed",
            "Form_Status": "Complete",
        })

        vitals.append({
            "Subject_ID": subject_id,
            "Visit_Name": visit_name,
            "Measurement_Date": planned_date,
            "Systolic_BP": random.randint(110, 138),
            "Diastolic_BP": random.randint(70, 88),
            "Heart_Rate": random.randint(62, 92),
            "Weight_kg": round(random.uniform(58, 96), 1),
        })

        # Illustrative values for testing data-management workflows only.
        # These values are not medical recommendations or clinical evidence.
        glucose = round(random.uniform(95, 175), 1)
        hba1c = round(random.uniform(6.5, 9.0), 1)
        creatinine = round(random.uniform(0.6, 1.3), 2)

        lab_values = [
            ("Glucose", glucose, "mg/dL"),
            ("HbA1c", hba1c, "%"),
            ("Creatinine", creatinine, "mg/dL"),
        ]

        for test_name, result, unit in lab_values:
            laboratory.append({
                "Lab_ID": f"LAB-{lab_counter:04d}",
                "Subject_ID": subject_id,
                "Visit_Name": visit_name,
                "Test_Name": test_name,
                "Result_Value": result,
                "Unit": unit,
                "Collection_Date": planned_date,
            })
            lab_counter += 1

        form_tracking.append({
            "Subject_ID": subject_id,
            "Visit_Name": visit_name,
            "Visit_Form_Status": "Complete",
            "Vitals_Form_Status": "Complete",
            "Laboratory_Form_Status": "Complete",
            "AE_Form_Status": "Complete",
            "Medication_Form_Status": "Complete",
            "Review_Status": "Reviewed",
        })

# -------------------------
# Adverse events (12 records)
# -------------------------
for i in range(12):
    subject_index = i * 2
    subject_id = subjects[subject_index]
    consent_date = demographics[subject_index]["Consent_Date"]
    onset_date = consent_date + timedelta(days=10 + i * 3)

    adverse_events.append({
        "AE_ID": f"AE-{i + 1:03d}",
        "Subject_ID": subject_id,
        "Event_Description": random.choice([
            "Headache",
            "Nausea",
            "Fatigue",
            "Dizziness",
            "Cough",
            "Abdominal discomfort",
        ]),
        "Onset_Date": onset_date,
        "Seriousness": "No",
        "Outcome": random.choice(["Resolved", "Recovering", "Ongoing"]),
    })

# -------------------------
# Concomitant medications (20 records)
# -------------------------
medication_names = [
    "Medication A",
    "Medication B",
    "Medication C",
    "Medication D",
]

for i in range(20):
    subject_index = i % SUBJECT_COUNT
    subject_id = subjects[subject_index]
    consent_date = demographics[subject_index]["Consent_Date"]
    med_start = consent_date + timedelta(days=2 + (i % 10))
    med_end = med_start + timedelta(days=15 + (i % 20))

    medications.append({
        "Medication_ID": f"MED-{i + 1:03d}",
        "Subject_ID": subject_id,
        "Medication_Name": medication_names[i % len(medication_names)],
        "Start_Date": med_start,
        "End_Date": med_end,
        "Indication": random.choice([
            "Medical history",
            "Concomitant treatment",
            "Other documented reason",
        ]),
    })


def write_workbook(path, sheets):
    """Write sheets with readable headers, filters, and Excel tables."""
    wb = Workbook()
    wb.remove(wb.active)

    header_fill = PatternFill("solid", fgColor="1F4E78")

    for sheet_name, records in sheets.items():
        ws = wb.create_sheet(sheet_name)

        if not records:
            ws.append(["No records"])
            continue

        headers = list(records[0].keys())
        ws.append(headers)

        for record in records:
            ws.append([record.get(header) for header in headers])

        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill
            cell.alignment = Alignment(wrap_text=True, vertical="top")

        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top")

                # Store date-only fields as Excel dates.
                if isinstance(cell.value, date):
                    cell.number_format = "yyyy-mm-dd"

        for col_cells in ws.columns:
            letter = col_cells[0].column_letter
            max_len = max(
                len(str(cell.value)) if cell.value is not None else 0
                for cell in col_cells
            )
            ws.column_dimensions[letter].width = min(max(max_len + 2, 14), 32)

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        ws.sheet_view.showGridLines = False

        table_ref = ws.dimensions
        if ws.max_row >= 2 and ws.max_column >= 1:
            safe_name = "".join(ch for ch in sheet_name if ch.isalnum())
            table = Table(
                displayName=f"Tbl{safe_name}",
                ref=table_ref,
            )
            table.tableStyleInfo = TableStyleInfo(
                name="TableStyleMedium2",
                showFirstColumn=False,
                showLastColumn=False,
                showRowStripes=True,
                showColumnStripes=False,
            )
            ws.add_table(table)

    wb.save(path)


clean_sheets = {
    "Demographics": demographics,
    "Eligibility": eligibility,
    "Visits": visits,
    "Vital_Signs": vitals,
    "Laboratory": laboratory,
    "Adverse_Events": adverse_events,
    "Medications": medications,
    "Form_Tracking": form_tracking,
}

clean_path = DATA_DIR / "clinical_study_clean.xlsx"
test_path = DATA_DIR / "clinical_study_test.xlsx"
defect_path = DOCS_DIR / "seeded_defect_register.xlsx"

write_workbook(clean_path, clean_sheets)

# Deep-copy the clean records before deliberately introducing defects.
test_sheets = deepcopy(clean_sheets)
defects = []


def register_defect(check_id, sheet, record_key, description):
    defects.append({
        "Defect_ID": f"DEF-{len(defects) + 1:03d}",
        "Related_Check_ID": check_id,
        "Sheet_Name": sheet,
        "Record_Reference": record_key,
        "Seeded_Defect": description,
        "Expected_Action": "Review the finding; document a query or resolution when needed.",
    })


# DEF-001: consent date missing while consent is Yes.
test_sheets["Demographics"][2]["Consent_Date"] = None
register_defect(
    "EC-004", "Demographics", "DM001-003",
    "Consent_Status is Yes but Consent_Date is missing.",
)

# DEF-002: completed visit without an actual date.
for record in test_sheets["Visits"]:
    if record["Subject_ID"] == "DM001-004" and record["Visit_Name"] == "Week 4":
        record["Actual_Date"] = None
        register_defect(
            "EC-009", "Visits", "DM001-004 / Week 4",
            "Visit is Completed but Actual_Date is missing.",
        )
        break

# DEF-003: duplicate scheduled visit record.
duplicate = next(
    record.copy()
    for record in test_sheets["Visits"]
    if record["Subject_ID"] == "DM001-005"
    and record["Visit_Name"] == "Week 8"
)
test_sheets["Visits"].append(duplicate)
register_defect(
    "EC-011", "Visits", "DM001-005 / Week 8",
    "Duplicate record for the same participant and scheduled visit.",
)

# DEF-004: missing lab result.
for record in test_sheets["Laboratory"]:
    if (
        record["Subject_ID"] == "DM001-006"
        and record["Visit_Name"] == "Week 4"
        and record["Test_Name"] == "HbA1c"
    ):
        record["Result_Value"] = None
        register_defect(
            "EC-017", "Laboratory", record["Lab_ID"],
            "Laboratory result value is missing.",
        )
        break

# DEF-005: lab unit does not match the test's defined unit.
for record in test_sheets["Laboratory"]:
    if (
        record["Subject_ID"] == "DM001-007"
        and record["Visit_Name"] == "Baseline"
        and record["Test_Name"] == "Glucose"
    ):
        record["Unit"] = "mmol/L"
        register_defect(
            "EC-018", "Laboratory", record["Lab_ID"],
            "Glucose unit differs from the project's specified unit.",
        )
        break

# DEF-006: AE onset before consent.
first_ae = test_sheets["Adverse_Events"][0]
subject = next(
    record for record in test_sheets["Demographics"]
    if record["Subject_ID"] == first_ae["Subject_ID"]
)
first_ae["Onset_Date"] = subject["Consent_Date"] - timedelta(days=1)
register_defect(
    "EC-022", "Adverse_Events", first_ae["AE_ID"],
    "Adverse event onset date precedes the recorded consent date.",
)

# DEF-007: medication end date precedes its start date.
first_med = test_sheets["Medications"][0]
first_med["End_Date"] = first_med["Start_Date"] - timedelta(days=1)
register_defect(
    "EC-025", "Medications", first_med["Medication_ID"],
    "Medication End_Date precedes Start_Date.",
)

# DEF-008: age outside the simulation's specified age range.
test_sheets["Demographics"][7]["Age"] = 17
register_defect(
    "EC-003", "Demographics", "DM001-008",
    "Age is outside the specified 18–70 training range; flag for review.",
)

write_workbook(test_path, test_sheets)
write_workbook(defect_path, {"Defect Register": defects})

print("Synthetic study datasets created successfully.")
print(f"Clean dataset: {clean_path}")
print(f"Test dataset:  {test_path}")
print(f"Defect register: {defect_path}")
print()
print("Clean dataset record counts:")
for name, records in clean_sheets.items():
    print(f"  {name}: {len(records)}")
print()
print(f"Seeded defects: {len(defects)}")
print("All records are fictional and generated for educational use.")
