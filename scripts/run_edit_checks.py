
from pathlib import Path
from collections import defaultdict
from datetime import date, datetime
from openpyxl import load_workbook, Workbook


# --------------------------------------------------
# Project paths
# --------------------------------------------------

BASE_DIR = Path(__file__).resolve().parent.parent
DATA_DIR = BASE_DIR / "data"
REPORT_DIR = BASE_DIR / "reports"

REPORT_DIR.mkdir(parents=True, exist_ok=True)


# --------------------------------------------------
# Helper functions
# --------------------------------------------------

def as_date(value):
    """Convert Excel date/datetime values into date objects."""
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return None


def is_missing(value):
    return value is None or (
        isinstance(value, str) and value.strip() == ""
    )


def read_sheet(workbook, sheet_name):
    """Read a worksheet into a list of dictionaries."""
    if sheet_name not in workbook.sheetnames:
        raise ValueError(
            f"Required sheet '{sheet_name}' not found. "
            f"Available sheets: {workbook.sheetnames}"
        )

    worksheet = workbook[sheet_name]
    rows = list(worksheet.iter_rows(values_only=True))

    if not rows:
        return []

    headers = [
        str(header).strip() if header is not None else ""
        for header in rows[0]
    ]

    records = []

    for values in rows[1:]:
        if all(is_missing(value) for value in values):
            continue

        record = {
            headers[index]: values[index]
            for index in range(len(headers))
            if headers[index]
        }
        records.append(record)

    return records


def add_finding(
    findings,
    check_id,
    subject_id,
    domain,
    record_reference,
    severity,
    description,
    requires_query
):
    findings.append({
        "Check_ID": check_id,
        "Subject_ID": subject_id,
        "Domain": domain,
        "Record_Reference": record_reference,
        "Severity": severity,
        "Finding_Description": description,
        "Requires_Query": requires_query,
    })


# --------------------------------------------------
# Edit-check engine
# --------------------------------------------------

def run_checks(input_file):
    print(f"\nReading: {input_file.name}")

    workbook = load_workbook(input_file, data_only=True)

    demographics = read_sheet(workbook, "Demographics")
    visits = read_sheet(workbook, "Visits")
    laboratory = read_sheet(workbook, "Laboratory")
    adverse_events = read_sheet(workbook, "Adverse_Events")
    medications = read_sheet(workbook, "Medications")

    findings = []

    # Build a subject-to-consent-date lookup.
    consent_dates = {}

    for row in demographics:
        subject_id = row.get("Subject_ID")
        consent_dates[subject_id] = as_date(
            row.get("Consent_Date")
        )

        # EC-004: Consent marked Yes but date missing.
        if (
            str(row.get("Consent_Status", "")).strip().lower()
            == "yes"
            and is_missing(row.get("Consent_Date"))
        ):
            add_finding(
                findings,
                "EC-004",
                subject_id,
                "Demographics",
                subject_id,
                "Error",
                "Consent status is Yes, but Consent_Date is missing.",
                "Yes"
            )

        # EC-003: Age outside the illustrative study range.
        age = row.get("Age")

        try:
            age_value = float(age)
        except (TypeError, ValueError):
            age_value = None

        if age_value is not None and not 18 <= age_value <= 70:
            add_finding(
                findings,
                "EC-003",
                subject_id,
                "Demographics",
                subject_id,
                "Review",
                f"Age {age_value:g} is outside the expected range of 18-70.",
                "No"
            )

    # EC-009: Completed visit without an actual date.
    for row in visits:
        status = str(row.get("Visit_Status", "")).strip().lower()

        if status == "completed" and is_missing(row.get("Actual_Date")):
            subject_id = row.get("Subject_ID")
            visit_name = row.get("Visit_Name")

            add_finding(
                findings,
                "EC-009",
                subject_id,
                "Visits",
                f"{subject_id} / {visit_name}",
                "Error",
                "Visit is marked Completed, but Actual_Date is missing.",
                "Yes"
            )

    # EC-011: Duplicate subject/visit combinations.
    visit_groups = defaultdict(list)

    for row in visits:
        key = (
            row.get("Subject_ID"),
            row.get("Visit_Name")
        )
        visit_groups[key].append(row)

    for (subject_id, visit_name), rows in visit_groups.items():
        if len(rows) > 1:
            add_finding(
                findings,
                "EC-011",
                subject_id,
                "Visits",
                f"{subject_id} / {visit_name}",
                "Error",
                f"Duplicate visit combination found ({len(rows)} rows).",
                "Yes"
            )

    # EC-017 and EC-018: Laboratory checks.
    expected_units = {
        "Glucose": "mg/dL",
        "HbA1c": "%",
        "Creatinine": "mg/dL",
    }

    for row in laboratory:
        subject_id = row.get("Subject_ID")
        lab_id = row.get("Lab_ID")
        test_name = str(row.get("Test_Name", "")).strip()
        result_value = row.get("Result_Value")
        unit = str(row.get("Unit", "")).strip()

        # EC-017: Missing, nonnumeric, or nonpositive result.
        try:
            numeric_result = float(result_value)
        except (TypeError, ValueError):
            numeric_result = None

        if numeric_result is None or numeric_result <= 0:
            add_finding(
                findings,
                "EC-017",
                subject_id,
                "Laboratory",
                lab_id,
                "Error",
                "Laboratory result is missing, nonnumeric, or not positive.",
                "Yes"
            )

        # EC-018: Unit inconsistent with the project specification.
        expected_unit = expected_units.get(test_name)

        if expected_unit is not None and unit != expected_unit:
            add_finding(
                findings,
                "EC-018",
                subject_id,
                "Laboratory",
                lab_id,
                "Error",
                f"{test_name} unit is '{unit}'; expected '{expected_unit}'.",
                "Yes"
            )

    # EC-022: Adverse event onset before consent.
    for row in adverse_events:
        subject_id = row.get("Subject_ID")
        ae_id = row.get("AE_ID")

        onset_date = as_date(row.get("Onset_Date"))
        consent_date = consent_dates.get(subject_id)

        if (
            onset_date is not None
            and consent_date is not None
            and onset_date < consent_date
        ):
            add_finding(
                findings,
                "EC-022",
                subject_id,
                "Adverse_Events",
                ae_id,
                "Review",
                "Adverse event onset date is before the consent date.",
                "Yes"
            )

    # EC-025: Medication end date before start date.
    for row in medications:
        subject_id = row.get("Subject_ID")
        medication_id = row.get("Medication_ID")

        start_date = as_date(row.get("Start_Date"))
        end_date = as_date(row.get("End_Date"))

        if (
            start_date is not None
            and end_date is not None
            and end_date < start_date
        ):
            add_finding(
                findings,
                "EC-025",
                subject_id,
                "Medications",
                medication_id,
                "Error",
                "Medication End_Date is earlier than Start_Date.",
                "Yes"
            )

    workbook.close()

    # Stable ordering and repeatable finding IDs.
    findings.sort(
        key=lambda item: (
            item["Check_ID"],
            str(item["Subject_ID"]),
            str(item["Record_Reference"])
        )
    )

    for index, finding in enumerate(findings, start=1):
        finding["Finding_ID"] = f"FND-{index:03d}"

    return findings


# --------------------------------------------------
# Findings report
# --------------------------------------------------

def write_report(findings, output_file, source_name):
    report = Workbook()

    summary = report.active
    summary.title = "Summary"

    error_count = sum(
        finding["Severity"] == "Error"
        for finding in findings
    )
    review_count = sum(
        finding["Severity"] == "Review"
        for finding in findings
    )

    summary.append(["Edit Check Report"])
    summary.append(["Source Dataset", source_name])
    summary.append(["Total Findings", len(findings)])
    summary.append(["Errors", error_count])
    summary.append(["Review Flags", review_count])
    summary.append([])
    summary.append([
        "Note",
        "Synthetic educational data only. Findings require review; "
        "they do not automatically change source data."
    ])

    details = report.create_sheet("Findings")

    columns = [
        "Finding_ID",
        "Check_ID",
        "Subject_ID",
        "Domain",
        "Record_Reference",
        "Severity",
        "Finding_Description",
        "Requires_Query",
    ]
    details.append(columns)

    for finding in findings:
        details.append([finding.get(column) for column in columns])

    for sheet in report.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions

        for column_cells in sheet.columns:
            max_length = max(
                len(str(cell.value or ""))
                for cell in column_cells
            )
            column_letter = column_cells[0].column_letter
            sheet.column_dimensions[column_letter].width = min(
                max(max_length + 2, 12),
                65
            )

    report.save(output_file)
    report.close()

    print(f"Total findings: {len(findings)}")
    print(f"Errors: {error_count}")
    print(f"Review flags: {review_count}")
    print(f"Report saved: {output_file}")


# --------------------------------------------------
# Run clean and test datasets separately
# --------------------------------------------------

if __name__ == "__main__":
    datasets = [
        (
            DATA_DIR / "clinical_study_clean.xlsx",
            REPORT_DIR / "clean_edit_check_findings.xlsx",
        ),
        (
            DATA_DIR / "clinical_study_test.xlsx",
            REPORT_DIR / "test_edit_check_findings.xlsx",
        ),
    ]

    for input_file, output_file in datasets:
        if not input_file.exists():
            raise FileNotFoundError(
                f"Dataset not found: {input_file}"
            )

        findings = run_checks(input_file)
        write_report(findings, output_file, input_file.name)

    print("\nEdit checks completed.")
